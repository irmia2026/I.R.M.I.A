"""XLSX 动作：建表 / 写值 / 校验 / 重算 / 读回 / 图表 / 条件格式 / 出图。

重算三条路，回执里**如实说用了哪条**：
  ① LibreOffice（随 DSH 打包的 libreoffice-kit `recalculate`）→ **保留公式且写回缓存值**，交付件就是算过的
  ② `formulas` 库 → 只在内存里算出数值，交付件**没有缓存值**（`delivery_file_recalculated: false`）
  ③ 都没有 → `deferred`，只把工作簿标成"打开时重算"（calcMode=auto），并明说"没算"

图表与条件格式是**原生 OOXML 对象**（不是往表里贴一张图）：用 openpyxl 的 chart / formatting
写进去，存完再重新打开数一遍 series 与规则条数——这是"真的建出来了"的判据。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import _common as C

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
VOLATILE = ("INDIRECT", "OFFSET", "NOW", "TODAY", "RAND", "RANDBETWEEN", "INFO", "CELL")
MODERN_ONLY = ("XLOOKUP", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "LET", "LAMBDA", "TEXTSPLIT", "XMATCH")

# 能建的图表（柱/横柱/线/饼/环/面积）。散点图要单独给 x 值，DSL 里说不清楚，先不做。
CHART_TYPES = ("bar", "barh", "line", "pie", "doughnut", "area")
CF_RULES = ("colorScale", "dataBar", "iconSet")
ICON_STYLES = ("3Arrows", "3ArrowsGray", "3Flags", "3TrafficLights1", "3TrafficLights2",
               "3Signs", "3Symbols", "3Symbols2", "4Arrows", "4ArrowsGray", "4RedToBlack",
               "4Rating", "4TrafficLights", "5Arrows", "5ArrowsGray", "5Rating", "5Quarters")


def _openpyxl():
    try:
        import openpyxl

        return openpyxl
    except Exception as exc:  # pragma: no cover
        raise C.OfficeError("DependencyMissing", f"需要 openpyxl：{exc}",
                            next_steps=["`python -m pip install openpyxl`"])


def _color(rgb: str | None):
    openpyxl = _openpyxl()
    if not rgb:
        return None
    rgb = rgb.strip().lstrip("#")
    if len(rgb) == 6:
        rgb = "FF" + rgb
    return openpyxl.styles.Color(rgb=rgb.upper())


def _apply_style(cell, spec: dict) -> None:
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    if spec.get("bold"):
        cell.font = Font(bold=True, color=_color(spec.get("font_color")) or _color("1A1F26"),
                         name=spec.get("font", "微软雅黑"), size=spec.get("size", 11))
    elif spec.get("size") or spec.get("font"):
        cell.font = Font(name=spec.get("font", "微软雅黑"), size=spec.get("size", 11),
                         color=_color(spec.get("font_color")))
    if spec.get("fill"):
        cell.fill = PatternFill("solid", fgColor=_color(spec["fill"]))
    if spec.get("number_format"):
        cell.number_format = spec["number_format"]
    if spec.get("wrap"):
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal=spec.get("align") or "left")
    elif spec.get("align"):
        cell.alignment = Alignment(horizontal=spec["align"], vertical="center")
    if spec.get("align") == "center" or spec.get("align") == "right":
        cell.alignment = Alignment(horizontal=spec["align"], vertical="center")
    if spec.get("border"):
        thin = Side(style="thin", color="FFC9D0D8")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _put(ws, ref: str, value, formula: bool = False) -> None:
    cell = ws[ref]
    if formula:
        text = str(value)
        cell.value = text if text.startswith("=") else "=" + text
    else:
        cell.value = value


# ── 图表 / 条件格式（原生 OOXML 对象） ──────────────────────────────────────


def _load_workbook(path: Path, data_only: bool = False):
    """load_workbook 的包装：**把 openpyxl 自己的 warning 收进回执**，不让它悄悄丢东西。

    openpyxl 遇到它不认的扩展块（比如条件格式的 x14、数据条的新外观）会警告
    "will be removed"——重新保存时确实会丢。这类信息必须出现在回执里，
    否则"加个图表"会把人家表里的条件格式外观弄回退而没人知道。
    """
    import warnings

    openpyxl = _openpyxl()
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(path, data_only=data_only)
    notes: list[str] = []
    for item in caught:
        msg = str(item.message).strip()
        if msg and msg not in notes:
            notes.append(msg)
    return wb, notes


def _ref_range(ws, rng: str):
    from openpyxl.chart import Reference
    from openpyxl.utils.cell import range_boundaries

    try:
        min_col, min_row, max_col, max_row = range_boundaries(str(rng).upper().replace("$", ""))
    except Exception as exc:
        raise C.OfficeError("BadSpec", f"区域写法不对：{rng!r}（要 A1:B4 这种）") from exc
    return Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)


def _strip_sheet(rng: str, sheet: str) -> str:
    """`汇总!B1:B4` → `B1:B4`；如果前缀是别的表就报错（别让人以为画的是那张表）。"""
    s = str(rng).strip()
    if "!" not in s:
        return s
    prefix, rest = s.rsplit("!", 1)
    prefix = prefix.strip("'").strip()
    if prefix and prefix.lower() != str(sheet).lower():
        raise C.OfficeError(
            "BadSpec",
            f"图表数据区域在「{prefix}」表，但目标是「{sheet}」表",
            next_steps=[f"用 --sheet {prefix} 指定那张表，或者把 --data 改成本表区域（如 B1:B4）"],
        )
    return rest


def _chart_title_text(ch) -> str:
    """把 openpyxl 的 Title 对象还原成字符串（用于"同标题就不重复加"的幂等判断）。"""
    t = getattr(ch, "title", None)
    if t is None:
        return ""
    try:
        return str(t.tx.rich.p[0].r[0].t)
    except Exception:
        try:
            return str(t.tx.strRef.f)
        except Exception:
            return str(t) if isinstance(t, str) else ""


def _build_chart_obj(ws, spec: dict):
    """按规格造一个 openpyxl 图表对象（不挂到表上）。"""
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart
    from openpyxl.chart.label import DataLabelList

    kind = str(spec.get("type") or "bar").lower()
    if kind not in CHART_TYPES:
        raise C.OfficeError(
            "BadSpec",
            f"不支持的图表类型：{kind!r}",
            next_steps=[f"支持：{' / '.join(CHART_TYPES)}（散点图暂不支持：需要单独的 x 值，DSL 里表达不清）"],
        )
    data = spec.get("data")
    if not data:
        raise C.OfficeError("BadSpec", "图表缺 data（数据区域，如 B1:B4）",
                            next_steps=['{"type":"bar","data":"B1:B4","categories":"A2:A4"}'])
    data = _strip_sheet(data, ws.title)
    titles_from_data = bool(spec.get("titles_from_data", True))
    title = spec.get("title")

    if kind == "bar":
        ch = BarChart()
        ch.type = "bar" if str(spec.get("direction") or "").lower() == "horizontal" else "col"
        grouping = str(spec.get("grouping") or "").lower()
        if grouping in ("stacked", "percentStacked"):
            ch.grouping = grouping
            ch.overlap = 100
    elif kind == "barh":
        ch = BarChart()
        ch.type = "bar"
    elif kind == "line":
        ch = LineChart()
    elif kind == "pie":
        ch = PieChart()
    elif kind == "doughnut":
        ch = DoughnutChart()
    else:
        ch = AreaChart()

    ch.title = title or None
    if spec.get("x_title"):
        ch.x_axis.title = str(spec["x_title"])
    if spec.get("y_title"):
        ch.y_axis.title = str(spec["y_title"])
    ch.width = float(spec.get("width") or 14)
    ch.height = float(spec.get("height") or 8)
    ch.add_data(_ref_range(ws, data), titles_from_data=titles_from_data)

    cats = spec.get("categories")
    if cats:
        ch.set_categories(_ref_range(ws, _strip_sheet(cats, ws.title)))
    elif kind in ("pie", "doughnut"):
        raise C.OfficeError(
            "BadSpec",
            "饼图/环图必须给 categories，否则切片没有标签（不是「看着能开」就算对）",
            next_steps=['{"type":"pie","data":"B1:B4","categories":"A2:A4"}'],
        )
    if kind in ("pie", "doughnut"):
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
    return ch


def _add_conditional(ws, spec: dict) -> dict:
    """给区域加一条条件格式（色阶 / 数据条 / 图标集）。"""
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule

    rng = spec.get("range")
    if not rng:
        raise C.OfficeError("BadSpec", "条件格式缺 range（如 B2:B4）")
    rule = str(spec.get("rule") or "colorScale")
    if rule not in CF_RULES:
        raise C.OfficeError("BadSpec", f"不支持的条件格式：{rule!r}",
                            next_steps=[f"支持：{' / '.join(CF_RULES)}"])
    if rule == "colorScale":
        colors = [str(c).lstrip("#") for c in (spec.get("colors") or ["F8696B", "FFEB84", "63BE7B"])]
        if len(colors) == 1:
            colors = [colors[0], "FFEB84", colors[0]]
        if len(colors) >= 3:
            obj = ColorScaleRule(start_type=spec.get("min_type", "min"), start_color=colors[0],
                                 mid_type="percentile", mid_value=50, mid_color=colors[1],
                                 end_type=spec.get("max_type", "max"), end_color=colors[2])
        else:
            obj = ColorScaleRule(start_type=spec.get("min_type", "min"), start_color=colors[0],
                                 end_type=spec.get("max_type", "max"), end_color=colors[-1])
    elif rule == "dataBar":
        obj = DataBarRule(start_type=spec.get("min_type", "min"), start_value=spec.get("min"),
                          end_type=spec.get("max_type", "max"), end_value=spec.get("max"),
                          color=str(spec.get("color") or "638EC6").lstrip("#").upper())
    else:
        icons = str(spec.get("icons") or "3Arrows")
        if icons not in ICON_STYLES:
            raise C.OfficeError("BadSpec", f"图标集名字不认识：{icons!r}",
                                next_steps=[f"可用的：{', '.join(ICON_STYLES[:6])} …"])
        reverse = bool(spec.get("reverse", False))
        obj = IconSetRule(icon_style=icons, type="percent",
                          values=spec.get("values") or [0, 33, 67], reverse=reverse)
    ws.conditional_formatting.add(str(rng), obj)
    return {"range": str(rng), "rule": rule}


def add_chart(xlsx_path: Path, out_xlsx: Path | None, spec: dict, sheet: str | None = None) -> dict:
    """往**已有**工作簿里加图表（同标题已存在就不重复加——重跑幂等）。"""
    out = out_xlsx or xlsx_path
    before = verify(xlsx_path)
    wb, load_warnings = _load_workbook(xlsx_path)
    name = sheet or wb.sheetnames[0]
    if name not in wb.sheetnames:
        raise C.OfficeError("SheetNotFound", f"没有这个工作表：{name}", next_steps=[f"现有：{wb.sheetnames}"])
    ws = wb[name]

    title = str(spec.get("title") or "")
    if title:
        existing = [c for c in ws._charts if _chart_title_text(c) == title]
        if existing:
            return {
                "out": C.rel_path(xlsx_path),
                "sheet": name,
                "added": False,
                "skipped_existing": True,
                "title": title,
                "charts": before["charts"],
                "conditional_rules": before["conditional_rules"],
                "chart_series_ok": before["chart_series_ok"],
                "preserved": {"conditional_rules_before": before["conditional_rules"],
                              "conditional_rules_after": before["conditional_rules"]},
                "warnings": [],
                "verify": _slim_verify(before),
                "note": "同标题图表已经在了，没有重复添加（重跑幂等）",
            }

    ch = _build_chart_obj(ws, spec)
    anchor = str(spec.get("anchor") or "E2")
    ws.add_chart(ch, anchor)
    C.prepare_write(out)
    wb.save(out)
    after = verify(out)
    warns = []
    dropped = before["conditional_rules"] - after["conditional_rules"]
    if dropped > 0:
        warns.append(
            f"加图表时整个工作簿被重写，条件格式少了 {dropped} 条"
            f"（{before['conditional_rules']} → {after['conditional_rules']}）"
        )
    if load_warnings:
        warns.append("openpyxl 读文件时丢掉了这些扩展块（重存后不在了）："
                     + "；".join(w[:120] for w in load_warnings[:2]))
    return {
        "out": C.rel_path(out),
        "sheet": name,
        "added": True,
        "anchor": anchor,
        "type": str(spec.get("type") or "bar"),
        "title": title or None,
        "charts": after["charts"],
        "conditional_rules": after["conditional_rules"],
        "chart_series_ok": after["chart_series_ok"],
        "preserved": {"conditional_rules_before": before["conditional_rules"],
                      "conditional_rules_after": after["conditional_rules"]},
        "warnings": warns,
        "verify": _slim_verify(after),
    }


def _slim_verify(report: dict) -> dict:
    return {k: report.get(k) for k in ("status", "formulas", "missing_cache_count", "charts",
                                       "conditional_rules", "chart_series_ok") if k in report}


def build(spec_path: Path, out_xlsx: Path) -> dict:
    openpyxl = _openpyxl()
    from openpyxl.utils import get_column_letter

    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    summary = []
    for sheet_spec in spec.get("sheets", []):
        name = str(sheet_spec.get("name") or "Sheet1")[:31]
        ws = wb.create_sheet(name)
        rows = sheet_spec.get("rows", [])
        start_row = int(sheet_spec.get("start_row", 1))
        start_col = int(sheet_spec.get("start_col", 1))
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                if value is None:
                    continue
                ws.cell(row=start_row + ri, column=start_col + ci, value=value)
        for ref, spec_val in (sheet_spec.get("formulas") or {}).items():
            _put(ws, ref, spec_val, formula=True)
        for ref, value in (sheet_spec.get("values") or {}).items():
            _put(ws, ref, value)
        for ref, style in (sheet_spec.get("styles") or {}).items():
            if ":" in ref:
                # 区间样式：逐个套
                for row in ws[ref]:
                    for cell in row if isinstance(row, tuple) else [row]:
                        _apply_style(cell, style)
            else:
                _apply_style(ws[ref], style)
        for col, width in (sheet_spec.get("cols") or {}).items():
            ws.column_dimensions[col].width = float(width)
        for row_idx, height in (sheet_spec.get("rows_height") or {}).items():
            ws.row_dimensions[int(row_idx)].height = float(height)
        for ref in sheet_spec.get("merges") or []:
            ws.merge_cells(ref)
        if sheet_spec.get("freeze"):
            ws.freeze_panes = sheet_spec["freeze"]
        charts_made = 0
        for chart_spec in sheet_spec.get("charts") or []:
            ch = _build_chart_obj(ws, chart_spec)
            ws.add_chart(ch, str(chart_spec.get("anchor") or "E2"))
            charts_made += 1
        cf_made = [_add_conditional(ws, cf) for cf in (sheet_spec.get("conditional") or [])]
        if sheet_spec.get("header_fill"):
            pass
        # 表头样式（第 1 行）
        if sheet_spec.get("header", True) and rows:
            from openpyxl.styles import Font, PatternFill

            for cell in ws[start_row]:
                if cell.value is None:
                    continue
                cell.font = Font(bold=True, name="微软雅黑", size=11, color="FF1A1F26")
                cell.fill = PatternFill("solid", fgColor="FFEDF1F6")
        ws.sheet_view.showGridLines = sheet_spec.get("gridlines", True)
        summary.append({"name": name, "rows": ws.max_row, "cols": ws.max_column,
                        "formulas": len(sheet_spec.get("formulas") or {}),
                        "charts": charts_made,
                        "conditional": len(cf_made)})

    if spec.get("calc_on_load", True):
        try:
            wb.calculation.fullCalcOnLoad = True
        except Exception:
            pass
    bak = C.prepare_write(out_xlsx)
    wb.save(out_xlsx)
    verify_res = verify(out_xlsx)
    return {
        "out": C.rel_path(out_xlsx),
        "bytes": out_xlsx.stat().st_size,
        "backup": C.rel_path(bak) if bak else None,
        "sheets": summary,
        "verify": verify_res,
    }


def set_cells(xlsx_path: Path, assignments: list[tuple[str, str, bool]], out_xlsx: Path | None = None) -> dict:
    """写单元值：assignments = [(sheet!A1, 值, 是否公式)]。"""
    openpyxl = _openpyxl()

    out = out_xlsx or xlsx_path
    wb = openpyxl.load_workbook(xlsx_path)
    changed = []
    for ref, value, is_formula in assignments:
        if "!" in ref:
            sheet_name, cell_ref = ref.rsplit("!", 1)
            sheet_name = sheet_name.strip("'")
            if sheet_name not in wb.sheetnames:
                raise C.OfficeError("SheetNotFound", f"没有这个工作表：{sheet_name}",
                                    next_steps=[f"现有：{wb.sheetnames}"])
            ws = wb[sheet_name]
        else:
            ws = wb.active
            cell_ref = ref
        before = ws[cell_ref].value
        _put(ws, cell_ref, value, formula=is_formula)
        changed.append({"cell": ref, "before": before, "after": ws[cell_ref].value})
    bak = C.prepare_write(out)
    wb.save(out)
    return {"out": C.rel_path(out), "backup": C.rel_path(bak) if bak else None,
            "changed": changed, "verify": verify(out)}


def _funcs_in(formula: str) -> list[str]:
    return [f.upper() for f in re.findall(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(", formula or "")]


def verify(xlsx_path: Path) -> dict:
    """静态体检：缺缓存值 / 错误值 / 禁用函数 / 隐式数组 / 结构风险 / 图表与条件格式。"""
    openpyxl = _openpyxl()

    errors: list[str] = []
    warnings: list[str] = []
    formulas_total = 0
    no_cache: list[str] = []
    error_cells: list[dict] = []
    forbidden: list[dict] = []
    volatile: list[dict] = []
    charts = 0
    charts_with_series = 0
    cf_rules = 0

    wb_f, notes_f = _load_workbook(xlsx_path, data_only=False)
    wb_v, notes_v = _load_workbook(xlsx_path, data_only=True)
    load_warnings = sorted(set(notes_f + notes_v))
    if len(wb_f.sheetnames) != len(set(wb_f.sheetnames)):
        errors.append("工作表重名")
    for ws in wb_f.worksheets:
        wsv = wb_v[ws.title] if ws.title in wb_v.sheetnames else None
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    formulas_total += 1
                    cached = wsv[cell.coordinate].value if wsv is not None else None
                    if cached is None:
                        no_cache.append(f"{ws.title}!{cell.coordinate}")
                    if isinstance(cached, str) and cached.upper() in ERROR_TOKENS:
                        error_cells.append({"cell": f"{ws.title}!{cell.coordinate}", "error": cached,
                                            "formula": val})
                    fns = _funcs_in(val)
                    for fn in fns:
                        if fn in MODERN_ONLY:
                            forbidden.append({"cell": f"{ws.title}!{cell.coordinate}", "function": fn,
                                              "why": "新版 Excel 才有，老版本/第三方库会算不出（#NAME?）"})
                        if fn in VOLATILE:
                            volatile.append({"cell": f"{ws.title}!{cell.coordinate}", "function": fn})
                    if val.startswith("={") or re.search(r"\{[^}]*,[^}]*\}", val):
                        warnings.append(f"{ws.title}!{cell.coordinate} 用了数组常量/隐式数组，缓存值语义依赖宿主")
                elif isinstance(val, str) and val.upper() in ERROR_TOKENS:
                    error_cells.append({"cell": f"{ws.title}!{cell.coordinate}", "error": val, "formula": None})
        # 合并单元格里放公式
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(row=rng.min_row, column=rng.min_col)
            if isinstance(top_left.value, str) and top_left.value.startswith("="):
                warnings.append(f"{ws.title}!{rng} 合并区左上角是公式，只有左上格算得对")
        # 图表 / 条件格式：存完重新打开数一遍（"真的建出来了"的判据）
        for ch in ws._charts:
            charts += 1
            if ch.series:
                charts_with_series += 1
            else:
                warnings.append(f"{ws.title} 里有图表没有 series（打开会是空白图）")
        for cf_range in ws.conditional_formatting:
            cf_rules += len(list(cf_range.rules))
    ext_links = list(wb_f._external_links) if hasattr(wb_f, "_external_links") else []
    if ext_links:
        warnings.append(f"有 {len(ext_links)} 个外部链接，换机器会断")

    if formulas_total and no_cache:
        status = "unverified"
    elif error_cells:
        status = "fail"
    elif forbidden or volatile or warnings:
        status = "pass_with_warnings"
    else:
        status = "pass"
    if no_cache:
        warnings.append(
            f"{len(no_cache)} 个公式没有缓存值：第三方阅读器/预览会显示空白或 0，"
            "交付前跑一次 `xlsx recalc`"
        )
    for item in forbidden:
        warnings.append(f"{item['cell']} 用了 {item['function']}：{item['why']}")
    if any("Conditional Formatting extension" in w for w in load_warnings):
        warnings.append(
            "这份工作簿里有 openpyxl 不认的条件格式扩展块（x14）：**重新保存会把外观丢掉**，"
            "所以别用 `xlsx chart` 去改这类文件；要加图就在 `xlsx build` 的 charts 里一次写进去"
        )
    return {
        "file": C.rel_path(xlsx_path),
        "status": status,
        "sheets": wb_f.sheetnames,
        "formulas": formulas_total,
        "missing_cache": no_cache[:50],
        "missing_cache_count": len(no_cache),
        "error_cells": error_cells[:20],
        "forbidden_functions": forbidden,
        "volatile_functions": volatile,
        "charts": charts,
        "chart_series_ok": (charts_with_series == charts),
        "conditional_rules": cf_rules,
        "openpyxl_warnings": load_warnings[:4],
        "errors": errors,
        "warnings": warnings,
    }


def read(xlsx_path: Path, sheet: str | None = None, cell_range: str | None = None) -> dict:
    openpyxl = _openpyxl()

    wb_f = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb_v = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = [sheet] if sheet else wb_f.sheetnames
    out = []
    for name in sheets:
        if name not in wb_f.sheetnames:
            raise C.OfficeError("SheetNotFound", f"没有这个工作表：{name}",
                                next_steps=[f"现有：{wb_f.sheetnames}"])
        ws = wb_f[name]
        wsv = wb_v[name] if name in wb_v.sheetnames else None
        target = ws[cell_range] if cell_range else ws
        rows = []
        if cell_range:
            for row in ws[cell_range]:
                rows.append([_cell_report(cell, wsv) for cell in (row if isinstance(row, tuple) else [row])])
        else:
            for row in ws.iter_rows():
                if all(c.value is None for c in row):
                    continue
                rows.append([_cell_report(cell, wsv) for cell in row])
        out.append({"name": name, "dims": ws.dimensions, "rows": rows})
    return {"file": C.rel_path(xlsx_path), "sheets": out}


def _cell_report(cell, wsv) -> dict:
    val = cell.value
    cached = None
    if isinstance(val, str) and val.startswith("=") and wsv is not None:
        cached = wsv[cell.coordinate].value
    item = {"ref": cell.coordinate, "value": _jsonable(val)}
    if isinstance(val, str) and val.startswith("="):
        item["formula"] = val
        item["cached"] = _jsonable(cached)
    if cell.number_format and cell.number_format != "General":
        item["fmt"] = cell.number_format
    return item


def _jsonable(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


# ── 重算 ────────────────────────────────────────────────────────────────────


def recalc(xlsx_path: Path, out_xlsx: Path | None = None, force_route: str | None = None) -> dict:
    """重算，返回"到底用哪条路、交付件有没有真的带上新值"。"""
    out = out_xlsx or xlsx_path.with_name(f"{xlsx_path.stem}_recalc{xlsx_path.suffix}")
    before = verify(xlsx_path)
    route = force_route
    detail: dict = {}

    if route in (None, "libreoffice"):
        import _lo

        info = _lo.discover()
        if info["ready"] and info["route"] in ("kit", "kit-env"):
            res = _lo.recalculate(xlsx_path, out)
            after = verify(out)
            detail = {
                "route": "libreoffice",
                "engine": "libreoffice-kit recalculate",
                "output": C.rel_path(out),
                "backup": res.get("backup"),
                "missing_fonts": res.get("missingFonts", []),
                "recalculated": True,
                "delivery_file_recalculated": True,
                "formulas_preserved": after["formulas"] == before["formulas"],
            }
            if after["missing_cache_count"] > 0:
                detail["warnings"] = [
                    f"重算后仍有 {after['missing_cache_count']} 个公式没缓存值（可能是跨表引用或名称未定义）"
                ]
            return {"status": "success", **detail, "before": _slim(before), "after": _slim(after)}
        if route == "libreoffice":
            raise C.OfficeError("DependencyMissing", "指定了 libreoffice 路由但没有可用的 LibreOffice",
                                next_steps=["去掉 --route libreoffice 让它自动降级"])

    if route in (None, "formulas"):
        try:
            import formulas
        except Exception:
            formulas = None
        if formulas is not None:
            return _recalc_with_formulas(xlsx_path, out, formulas, before)

    # ③ 降级：只标"打开时重算"
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    C.prepare_write(out)
    wb.save(out)
    return {
        "status": "deferred",
        "route": "calc_mode_auto",
        "recalculated": False,
        "delivery_file_recalculated": False,
        "output": C.rel_path(out),
        "note": "没有 LibreOffice 也没有 formulas：**没有算**，只是让 Excel 打开时自己重算",
        "next": [
            "`python -m pip install formulas`（纯 Python，能算出值但不写回交付件）",
            "或者让 LibreOffice 可用（随 DSH 打包的 libreoffice-kit 就行），那条能真正写回缓存",
        ],
        "before": _slim(before),
    }


def _recalc_with_formulas(xlsx_path: Path, out: Path, formulas_mod, before: dict) -> dict:
    """formulas 库：算出数值，但**交付件不写缓存**——这一点必须在回执里说清楚。"""
    import shutil

    work = out.with_name(f"{out.stem}_formulas_verify{out.suffix}")
    shutil.copy2(xlsx_path, work)
    computed: dict = {}
    status = "success"
    error = None
    try:
        # formulas 会往 stderr 打 tqdm 进度条，--json 回执要干净，先关掉它
        import contextlib
        import io

        with contextlib.redirect_stderr(io.StringIO()):
            model = formulas_mod.ExcelModel().loads(str(xlsx_path)).finish()
            solution = model.calculate()
        for key, value in solution.items():
            try:
                val = value.value[0, 0] if hasattr(value, "value") else value
                computed[str(key)] = _jsonable(val)
            except Exception:
                continue
    except Exception as exc:  # noqa: BLE001
        status = "partial"
        error = f"{type(exc).__name__}: {exc}"
    # 交付件照抄一份（缓存仍是空的），只把公式标成打开时重算
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    C.prepare_write(out)
    wb.save(out)
    return {
        "status": status,
        "route": "formulas",
        "recalculated": bool(computed),
        "delivery_file_recalculated": False,
        "output": C.rel_path(out),
        "verify_copy": C.rel_path(work),
        "computed_count": len(computed),
        "computed": dict(list(computed.items())[:60]),
        "error": error,
        "note": (
            "formulas 只把**数值算出来放在回执里**，交付的 xlsx 里缓存值仍然是空的"
            "（openpyxl 不支持把公式与缓存同时写回；硬写会把图表/样式弄坏）。"
            "要让交付件自带新值，得走 LibreOffice 那条。"
        ),
        "before": _slim(before),
    }


def _slim(report: dict) -> dict:
    return {k: report.get(k) for k in ("status", "formulas", "missing_cache_count", "error_cells",
                                       "forbidden_functions") if k in report}


# ── 出图 ────────────────────────────────────────────────────────────────────


def render(xlsx_path: Path, out_dir: Path, dpi: int = 130, sheet: str | None = None,
           cell_range: str | None = None, engine: str = "pdf") -> dict:
    """工作表出图。

    默认走 **LibreOffice 转 PDF → pdfium 出图**（可信、不裁切）。
    只有显式要 `--sheet/--range` 时才用 libreoffice-kit 的裁到数据区渲染，
    并把"实测有 4/3 放大裁切"写进 degraded —— 那张图只能当参考。
    """
    import _lo
    import _render

    out_dir.mkdir(parents=True, exist_ok=True)
    degraded: list[dict] = []
    images: list[dict] = []
    missing_fonts: list[str] = []

    if sheet or cell_range or engine == "kit":
        openpyxl = _openpyxl()
        wb = openpyxl.load_workbook(xlsx_path)
        names = [sheet] if sheet else list(wb.sheetnames)
        for i, name in enumerate(names):
            part_dir = out_dir / f"_kit-{i}"
            if part_dir.exists():
                import shutil

                shutil.rmtree(part_dir)
            res = _lo.render(xlsx_path, part_dir, dpi=dpi, sheet=name, cell_range=cell_range)
            missing_fonts.extend(res.get("missingFonts") or [])
            for img in res["images"]:
                dst = out_dir / f"page-{len(images) + 1:03d}.png"
                C.prepare_write(dst, backup=False)
                import shutil as _sh

                _sh.copy2(img["path"], dst)
                images.append({"index": len(images) + 1, "page": len(images) + 1, "sheet": name,
                               "range": img.get("range"), "path": str(dst),
                               "width": img.get("width"), "height": img.get("height")})
        route = "libreoffice-kit（裁到数据区）"
        degraded.append({
            "capability": "按 sheet/range 精确裁图",
            "reason": "实测 libreoffice-kit 的 render 页面按 dpi/72 出像素、内容按 dpi/96 排版，"
                      "图右侧会被裁掉约 1/4；这张图只能当参考，不能当「版式对不对」的证据",
            "fallback": "去掉 --sheet/--range 让默认的 convert→pdf→png 路线出整页图（可信）",
        })
    else:
        res = _render.render_office_via_pdf(xlsx_path, out_dir, dpi=dpi)
        images = res["images"]
        route = "libreoffice→pdf→pdfium"

    if not images:
        raise C.OfficeError("RenderEmpty", "没吐出工作表图",
                            next_steps=["确认工作簿里有可见数据（空表不渲染）"])
    stats = _render.summarize_pages(images)
    labels = [(i.get("sheet") or f"第 {i.get('page')} 页") for i in images]
    cs = _render.contact_sheet(images, out_dir / "contact-sheet.png", cols=2,
                               title=f"{xlsx_path.name} · {len(images)} 张工作表图 · {dpi}dpi",
                               labels=labels)
    return {
        "file": C.rel_path(xlsx_path),
        "route": route,
        "dpi": dpi,
        "image_count": len(images),
        "images": [{**i, "path": C.rel_path(i["path"])} for i in images],
        "contact_sheet": C.rel_path(cs["path"]),
        "missing_fonts": sorted(set(missing_fonts)),
        "pages": [{**s, "path": C.rel_path(s["path"])} for s in stats],
        "blank_pages": [s["page"] for s in stats if s["blank"]],
        "mean_ink_ratio": round(sum(s["ink_ratio"] for s in stats) / len(stats), 5) if stats else 0.0,
        "degraded": degraded,
        "note": "表格出图是「打印出来长什么样」；列宽/换行请回 Excel 里看",
    }


def _sheet_of(img: dict) -> str:
    return f"{img.get('sheet') or ''}".strip() or "工作表"
